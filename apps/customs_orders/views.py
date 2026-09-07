from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from .models import CustomsOrder
from .services import (
    CustomsOrderError,
    create_customs_order_from_boundary,
    eligible_customs_sources,
)


@login_required
def customs_orders_list(request):
    if not request.user.can_view_purchase_cost:
        raise PermissionDenied
    return render(request, "customs_orders/list.html", {"orders": CustomsOrder.objects.all()})


@login_required
def customs_order_detail(request, pk):
    if not request.user.can_view_purchase_cost:
        raise PermissionDenied
    order = get_object_or_404(CustomsOrder.objects.prefetch_related("lines"), pk=pk)
    if request.GET.get("export") == "1":
        import openpyxl

        from apps.actions.services import export_customs_xlsx

        def rows(lines):
            return [{
                "number": line.article, "name_ru": line.name_ru, "name_en": line.name_en,
                "manufacturer": line.manufacturer, "country": "", "gross_weight_kg": None,
                "net_weight_kg": None, "quantity": line.quantity, "usd_price": line.wholesale_usd,
                "application_area": "",
                "provenance": "ordered" if line.is_ordered else "sales_repairs",
            } for line in lines]
        originals = [line for line in order.lines.all() if not line.is_analog]
        analogs = [line for line in order.lines.all() if line.is_analog]
        source = openpyxl.load_workbook(export_customs_xlsx(rows=rows(originals)))
        analog_book = openpyxl.load_workbook(export_customs_xlsx(rows=rows(analogs)))
        source.active.title = "Оригиналы"
        target = source.create_sheet("Аналоги")
        for row in analog_book.active.iter_rows():
            for cell in row:
                target.cell(cell.row, cell.column, cell.value)
        import io
        output = io.BytesIO()
        source.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="customs_order_{order.number}.xlsx"'
        )
        return response
    return render(request, "customs_orders/detail.html", {"order": order})


@login_required
def customs_order_selection(request):
    if not request.user.can_view_purchase_cost:
        raise PermissionDenied
    sources = eligible_customs_sources()
    if request.method == "POST":
        raw = request.POST.get("boundary", "")
        try:
            kind, source_id = raw.split(":", 1)
            order = create_customs_order_from_boundary(
                number=int(request.POST.get("number", "0")),
                boundary_source=(kind, int(source_id)), by=request.user,
            )
        except (ValueError, CustomsOrderError) as exc:
            messages.error(request, str(exc))
        else:
            return redirect("customs_order_detail", pk=order.pk)
    return render(request, "customs_orders/select.html", {
        "sources": sources, "initialization_required": not CustomsOrder.objects.exists(),
    })
