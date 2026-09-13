"""Import Arctic Cat dealer catalog facts through the shared import workflow.

The adapter accepts the documented ``usprice`` sheet only.  Arctic Cat dealer
price is source data, not a confirmed DenisStock customer-price policy, so it
never recalculates ``PartType.recommended_price``.  Package quantity remains
supplier metadata and never creates stock.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from django.db.models import Max
from django.utils import timezone

from apps.catalog.models import Category, Manufacturer, PartNumber, PartType, Unit, normalize_number

from .models import ArcticCatCatalogPart

MAX_PROBLEMS = 200
MAX_TEXT = 10_000
SOURCE = ArcticCatCatalogPart.SOURCE_DEALER
ARCTIC_CAT_MANUFACTURER = "Arctic Cat"
ARCTIC_CAT_CATEGORY = "Arctic Cat catalog"
DEFAULT_UNIT_NAME = "Штука"
SHEET_NAME = "usprice"
FORMAT = "ARCTIC_CAT_DEALER_CATALOG"

# This intentionally recognizes only the documented whole-cell reference.
# A prose description beginning with R/B is not a supersession instruction.
REPLACEMENT_RE = re.compile(r"^R/B\s+([0-9]{4}-[0-9]{3})$", flags=re.IGNORECASE)

HEADER_ALIASES = {
    "article": {"p/n", "pn", "part number", "part no", "part #"},
    "description": {"description", "part description"},
    "package_quantity": {"pkg qty", "package qty", "package quantity"},
    "dealer_price": {"dealer price", "dealer price usd"},
}
REQUIRED_HEADERS = ("article", "description", "package_quantity", "dealer_price")


class ArcticCatCatalogError(RuntimeError):
    """An Arctic Cat workbook is not safe to interpret."""


@dataclass(frozen=True)
class Problem:
    row: int
    reason: str
    detail: str = ""


@dataclass(frozen=True)
class IncomingRow:
    number: int
    article: str
    normalized_article: str
    description: str
    package_quantity: str
    dealer_price_usd: Decimal | None
    dealer_price_state: str
    replacement_article: str
    normalized_replacement_article: str

    @property
    def identity(self) -> str:
        return self.normalized_article


@dataclass
class Plan:
    sheet: str = ""
    rows_scanned: int = 0
    blank_rows: int = 0
    valid: int = 0
    unique_part_numbers: int = 0
    new: int = 0
    existing: int = 0
    unchanged: int = 0
    description_changed: int = 0
    price_changed: int = 0
    replacement_changed: int = 0
    package_quantity_changed: int = 0
    zero_price_rows: int = 0
    blank_price_rows: int = 0
    duplicate_part_numbers: int = 0
    warnings: int = 0
    errors: int = 0
    problems: list[Problem] = field(default_factory=list)

    def problem(self, row: int, reason: str, detail: str = "", *, error: bool) -> None:
        if error:
            self.errors += 1
        else:
            self.warnings += 1
        if len(self.problems) < MAX_PROBLEMS:
            self.problems.append(Problem(row, reason, detail))

    def as_summary(self) -> dict:
        return {
            "format": FORMAT,
            "format_label": "Arctic Cat dealer catalog",
            "sheet": self.sheet,
            "rows_scanned": self.rows_scanned,
            "blank_rows": self.blank_rows,
            "valid": self.valid,
            "unique_part_numbers": self.unique_part_numbers,
            "new": self.new,
            "existing": self.existing,
            "unchanged": self.unchanged,
            "description_changed": self.description_changed,
            "price_changed": self.price_changed,
            "replacement_changed": self.replacement_changed,
            "package_quantity_changed": self.package_quantity_changed,
            "zero_price_rows": self.zero_price_rows,
            "blank_price_rows": self.blank_price_rows,
            "duplicate_part_numbers": self.duplicate_part_numbers,
            "warnings": self.warnings,
            "errors": self.errors,
            "currency": "USD",
            "price_policy": "raw_supplier_only",
            "stock_changes": False,
            "problems": [item.__dict__ for item in self.problems],
            "problems_total": self.warnings + self.errors,
        }


def _header_key(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").strip().lower().split())


def _text(value: object) -> str:
    value = "" if value is None else str(value).strip()
    if len(value) > MAX_TEXT:
        raise ArcticCatCatalogError("В ячейке слишком длинное значение.")
    return value


def _identifier(cell) -> str:
    """Keep text cells exact and recover a simple zero-padded numeric cell."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        number_format = str(cell.number_format or "")
        if re.fullmatch(r"0+(?:-0+)*", number_format):
            width = number_format.count("0")
            digits = iter(f"{value:0{width}d}")
            return "".join(next(digits) if char == "0" else char for char in number_format)
    return _text(value)


def _dealer_price(value: object) -> tuple[Decimal | None, str]:
    text = _text(value).replace(" ", "").replace(",", ".")
    if not text:
        return None, ArcticCatCatalogPart.DealerPriceState.BLANK
    try:
        price = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ArcticCatCatalogError(f"DEALER PRICE: «{text}» не является ценой.") from exc
    if not price.is_finite() or price < 0:
        raise ArcticCatCatalogError("DEALER PRICE: цена должна быть конечной и неотрицательной.")
    try:
        price = price.quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ArcticCatCatalogError("DEALER PRICE: цена слишком велика.") from exc
    if price == 0:
        return None, ArcticCatCatalogPart.DealerPriceState.ZERO
    return price, ArcticCatCatalogPart.DealerPriceState.KNOWN


def _sheet_name(names) -> str:
    matches = [name for name in names if _header_key(name) == SHEET_NAME]
    if len(matches) == 1:
        return matches[0]
    found = ", ".join(names) or "ни одного"
    if not matches:
        raise ArcticCatCatalogError(f"Для Arctic Cat нужен лист usprice. Найдены листы: {found}")
    raise ArcticCatCatalogError(f"В книге несколько листов usprice: {', '.join(matches)}")


def _mapping(cells) -> dict[str, int]:
    found: dict[str, int] = {}
    actual = []
    for index, cell in enumerate(cells):
        key = _header_key(cell.value)
        actual.append(_text(cell.value))
        for name, aliases in HEADER_ALIASES.items():
            if key in aliases and name not in found:
                found[name] = index
                break
    missing = [name for name in REQUIRED_HEADERS if name not in found]
    if missing:
        expected = ", ".join(("P/N", "Description", "Pkg Qty", "DEALER PRICE"))
        got = ", ".join(value for value in actual if value) or "пусто"
        raise ArcticCatCatalogError(
            "Не распознан Arctic Cat-каталог. Ожидаются заголовки "
            f"{expected}; отсутствуют: {', '.join(missing)}. Найдены: {got}"
        )
    return found


def _replacement(description: str) -> tuple[str, str, bool]:
    match = REPLACEMENT_RE.fullmatch(description)
    if match:
        article = " ".join(match.group(1).split())
        return article, normalize_number(article), False
    # A near-match remains a human-readable source description, not a guessed link.
    return "", "", bool(re.match(r"^R/B\b", description, flags=re.IGNORECASE))


def _read(path: Path) -> tuple[list[IncomingRow], Plan]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ArcticCatCatalogError("Не установлен openpyxl.") from exc
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise ArcticCatCatalogError("Ожидается доступный файл .xlsx.")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ArcticCatCatalogError(f"Файл не читается как Excel: {exc}") from exc
    try:
        worksheet = workbook[_sheet_name(workbook.sheetnames)]
        rows = worksheet.iter_rows()
        try:
            mapping = _mapping(next(rows))
        except StopIteration as exc:
            raise ArcticCatCatalogError("Файл пустой.") from exc
        plan = Plan(sheet=worksheet.title)
        records: list[IncomingRow] = []
        seen: dict[str, IncomingRow] = {}
        for row_number, cells in enumerate(rows, start=2):
            values = [_text(cell.value) for cell in cells]
            if not any(values):
                plan.blank_rows += 1
                continue
            plan.rows_scanned += 1
            mapped = {name: cells[index] for name, index in mapping.items() if index < len(cells)}
            article = _identifier(mapped["article"])[:100]
            description = " ".join(_text(mapped["description"].value).split())[:255]
            package_quantity = _identifier(mapped["package_quantity"])[:80]
            if not article or not description:
                plan.problem(row_number, "Не заполнены обязательные поля", error=True)
                continue
            normalized = normalize_number(article)
            if not normalized:
                plan.problem(row_number, "Некорректный P/N", article, error=True)
                continue
            try:
                price, price_state = _dealer_price(mapped["dealer_price"].value)
            except ArcticCatCatalogError as exc:
                plan.problem(row_number, "Некорректная цена", str(exc), error=True)
                continue
            if price_state == ArcticCatCatalogPart.DealerPriceState.ZERO:
                plan.zero_price_rows += 1
            elif price_state == ArcticCatCatalogPart.DealerPriceState.BLANK:
                plan.blank_price_rows += 1
            replacement_article, normalized_replacement, replacement_warning = _replacement(
                description
            )
            if replacement_warning:
                plan.problem(
                    row_number,
                    "Неоднозначная R/B-запись не стала заменой",
                    description,
                    error=False,
                )
            record = IncomingRow(
                number=row_number,
                article=article,
                normalized_article=normalized,
                description=description,
                package_quantity=package_quantity,
                dealer_price_usd=price,
                dealer_price_state=price_state,
                replacement_article=replacement_article,
                normalized_replacement_article=normalized_replacement,
            )
            previous = seen.get(record.identity)
            if previous is not None:
                plan.duplicate_part_numbers += 1
                if previous != record:
                    plan.problem(
                        row_number,
                        "Конфликтующий повтор P/N",
                        record.article,
                        error=True,
                    )
                else:
                    plan.problem(row_number, "Повтор P/N", record.article, error=False)
                continue
            seen[record.identity] = record
            records.append(record)
        plan.unique_part_numbers = len(records)
        return records, plan
    finally:
        if workbook is not None:
            workbook.close()


def _source_index(records: list[IncomingRow]) -> dict[str, ArcticCatCatalogPart]:
    return {
        entry.normalized_supplier_article: entry
        for entry in ArcticCatCatalogPart.objects.filter(
            source=SOURCE,
            normalized_supplier_article__in=[record.normalized_article for record in records],
        ).select_related("part")
    }


def _stored_price(entry: ArcticCatCatalogPart, record: IncomingRow) -> tuple[Decimal | None, str]:
    """Never let an unavailable price erase a previously known positive value."""
    if record.dealer_price_state == ArcticCatCatalogPart.DealerPriceState.KNOWN:
        return record.dealer_price_usd, record.dealer_price_state
    if entry.dealer_price_usd is not None and entry.dealer_price_usd > 0:
        return entry.dealer_price_usd, entry.dealer_price_state
    return None, record.dealer_price_state


def _changes(entry: ArcticCatCatalogPart, record: IncomingRow) -> set[str]:
    stored_price, stored_state = _stored_price(entry, record)
    changes = set()
    if entry.source_description != record.description:
        changes.add("description")
    if entry.package_quantity != record.package_quantity:
        changes.add("package_quantity")
    if (entry.dealer_price_usd, entry.dealer_price_state) != (stored_price, stored_state):
        changes.add("price")
    if (
        entry.replacement_article,
        entry.normalized_replacement_article,
    ) != (record.replacement_article, record.normalized_replacement_article):
        changes.add("replacement")
    return changes


def _classify(records: list[IncomingRow], plan: Plan) -> None:
    entries = _source_index(records)
    for record in records:
        entry = entries.get(record.identity)
        plan.valid += 1
        if entry is None:
            plan.new += 1
            continue
        plan.existing += 1
        changes = _changes(entry, record)
        if not changes:
            plan.unchanged += 1
            continue
        plan.description_changed += "description" in changes
        plan.package_quantity_changed += "package_quantity" in changes
        plan.price_changed += "price" in changes
        plan.replacement_changed += "replacement" in changes


def build_plan(path) -> Plan:
    records, plan = _read(Path(path))
    _classify(records, plan)
    return plan


def catalog_fingerprint() -> str:
    aggregate = ArcticCatCatalogPart.objects.aggregate(total=Max("pk"), touched=Max("updated_at"))
    payload = f"{ArcticCatCatalogPart.objects.count()}|{aggregate['total']}|{aggregate['touched']}"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _unit() -> Unit:
    unit = Unit.objects.filter(name__iexact=DEFAULT_UNIT_NAME, is_active=True).first()
    if unit is None:
        unit = Unit.objects.filter(is_active=True).first()
    if unit is None:
        raise ArcticCatCatalogError("В справочниках нет активной единицы измерения.")
    return unit


@transaction.atomic
def apply_file(path) -> dict:
    records, plan = _read(Path(path))
    _classify(records, plan)
    entries = _source_index(records)
    category, _ = Category.objects.get_or_create(name=ARCTIC_CAT_CATEGORY, parent=None)
    manufacturer, _ = Manufacturer.objects.get_or_create(name=ARCTIC_CAT_MANUFACTURER)
    unit = _unit()
    to_create = [record for record in records if record.identity not in entries]
    cards = [
        PartType(
            name=record.description,
            description=record.description,
            category=category,
            manufacturer=manufacturer,
            unit=unit,
            tracking_mode=PartType.TrackingMode.BULK,
            # Supplier dealer price is deliberately not a customer-price policy.
            recommended_price=None,
            # Public exposure remains a separate explicit eligibility decision.
            is_public=False,
        )
        for record in to_create
    ]
    if cards:
        PartType.objects.bulk_create(cards, batch_size=1000)
        PartNumber.objects.bulk_create(
            [
                PartNumber(
                    part=part,
                    value=record.article,
                    normalized_value=record.normalized_article,
                    kind=PartNumber.Kind.ARTICLE,
                    is_primary=True,
                )
                for record, part in zip(to_create, cards, strict=True)
            ],
            batch_size=1000,
        )
        ArcticCatCatalogPart.objects.bulk_create(
            [
                ArcticCatCatalogPart(
                    source=SOURCE,
                    part=part,
                    supplier_article=record.article,
                    normalized_supplier_article=record.normalized_article,
                    source_description=record.description,
                    package_quantity=record.package_quantity,
                    dealer_price_usd=record.dealer_price_usd,
                    dealer_price_state=record.dealer_price_state,
                    replacement_article=record.replacement_article,
                    normalized_replacement_article=record.normalized_replacement_article,
                )
                for record, part in zip(to_create, cards, strict=True)
            ],
            batch_size=1000,
        )

    updated_entries: list[ArcticCatCatalogPart] = []
    updated_parts: list[PartType] = []
    for record in records:
        entry = entries.get(record.identity)
        if entry is None:
            continue
        changes = _changes(entry, record)
        if not changes:
            continue
        price, price_state = _stored_price(entry, record)
        entry.source_description = record.description
        entry.package_quantity = record.package_quantity
        entry.dealer_price_usd = price
        entry.dealer_price_state = price_state
        entry.replacement_article = record.replacement_article
        entry.normalized_replacement_article = record.normalized_replacement_article
        entry.updated_at = timezone.now()
        updated_entries.append(entry)
        if "description" in changes:
            entry.part.name = record.description
            entry.part.description = record.description
            updated_parts.append(entry.part)
    if updated_parts:
        PartType.objects.bulk_update(updated_parts, ["name", "description"], batch_size=1000)
    if updated_entries:
        ArcticCatCatalogPart.objects.bulk_update(
            updated_entries,
            [
                "source_description",
                "package_quantity",
                "dealer_price_usd",
                "dealer_price_state",
                "replacement_article",
                "normalized_replacement_article",
                "updated_at",
            ],
            batch_size=1000,
        )
    result = plan.as_summary()
    result.update(
        {
            "created_parts": len(to_create),
            "updated_parts": len(updated_entries),
            "stock_changes": False,
        }
    )
    return result
