"""Сохранённый заказ выгружается целиком с неизменными данными и оформлением."""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse

from apps.actions.services import TEMPLATE_PATH
from apps.brp.models import BrpCatalogPart
from apps.catalog.models import Category, Manufacturer, PartType, Unit
from apps.customs_orders.export import export_customs_order_xlsx
from apps.customs_orders.models import CustomsOrder, CustomsOrderLine
from apps.warehouse.models import ValuationSettings

pytestmark = pytest.mark.django_db


@pytest.fixture
def order():
    return CustomsOrder.objects.create(
        number=125, fx_rate=Decimal("97.1250"),
        total_quantity=Decimal("6.500"), total_rub=Decimal("37498.90"),
    )


def _line(order, source_id=1, **changes):
    values = {
        "source": CustomsOrderLine.Source.SALE,
        "source_id": source_id,
        "article": "420001234",
        "name_ru": "ПОДШИПНИК",
        "name_en": "BALL BEARING",
        "manufacturer": "BRP",
        "country": "CANADA",
        "gross_weight_kg": None,
        "net_weight_kg": None,
        "application_area": "СНЕГОХОД",
        "quantity": Decimal("2.500"),
        "wholesale_usd": Decimal("10.2500"),
        "rub_amount": Decimal("2488.83"),
    }
    values.update(changes)
    return CustomsOrderLine.objects.create(order=order, **values)


def _article_rows(sheet):
    return [row for row in range(10, sheet.max_row) if sheet[f"B{row}"].value is not None]


def _export_rows(order, sheet_name="Оригиналы"):
    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))[sheet_name]
    return [
        tuple(sheet[f"{column}{row}"].value for column in "BCDEFGHIJKLM")
        for row in _article_rows(sheet)
    ]


def _occurred(day, hour=12):
    return datetime(2026, 9, day, hour, tzinfo=UTC)


def _part(name, manufacturer):
    category = Category.objects.create(name=f"customs-{name}")
    unit = Unit.objects.get(name="Штука")
    maker = Manufacturer.objects.get_or_create(name=manufacturer)[0]
    return PartType.objects.create(
        name=name, category=category, unit=unit, manufacturer=maker,
        tracking_mode=PartType.TrackingMode.BULK,
    )


def test_order_export_fills_missing_names_from_stable_catalog_link(order):
    from apps.brp.models import BrpPartLink

    part = _part("Карточка детали", "BRP")
    catalog = BrpCatalogPart.objects.create(
        material_no="001.5600.000", part_desc="OIL SEAL", is_current=True,
    )
    BrpPartLink.objects.create(
        part=part, brp_part=catalog, usd_rate_used=Decimal("100"),
        markup_percent_used=Decimal("0"),
    )
    _line(
        order, article="001.5600.000", manufacturer="BRP", name_ru="", name_en="",
        part_type=part,
    )
    _line(
        order, 2, article="001.5600.000", manufacturer="BRP", name_ru="", name_en="",
        part_type=part, is_analog=True,
    )

    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    sheet = book["Оригиналы"]
    assert sheet["B10"].value == "001.5600.000"
    assert sheet["C10"].value == "САЛЬНИК"
    assert sheet["D10"].value == "OIL SEAL"
    assert book["Аналоги"]["C10"].value == "САЛЬНИК"
    assert book["Аналоги"]["D10"].value == "OIL SEAL"


def test_order_export_uses_historical_customs_version_before_current_catalog():
    from apps.actions.models import PartCustomsDataVersion

    order = CustomsOrder.objects.create(
        number=127, fx_rate=Decimal("105.0000"), total_quantity=Decimal("1"),
        total_rub=Decimal("1050.00"),
    )
    part = _part("Historical card", "BRP")
    PartCustomsDataVersion.objects.create(
        part_type=part, version=1, effective_from=_occurred(1),
        customs_name_ru="ИСТОРИЧЕСКОЕ ИМЯ", customs_name_ru_confirmed=True,
        customs_name_en="HISTORICAL NAME",
    )
    PartCustomsDataVersion.objects.create(
        part_type=part, version=2, effective_from=_occurred(20),
        customs_name_ru="НОВОЕ ИМЯ", customs_name_ru_confirmed=True,
        customs_name_en="NEW NAME",
    )
    BrpCatalogPart.objects.create(material_no="HIST-1", part_desc="CURRENT CATALOG NAME")
    _line(order, article="HIST-1", manufacturer="BRP", name_ru="", name_en="",
          part_type=part, occurred_at=_occurred(10))

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert sheet["C10"].value == "ИСТОРИЧЕСКОЕ ИМЯ"
    assert sheet["D10"].value == "HISTORICAL NAME"


def test_order_export_exact_article_fallback_is_manufacturer_qualified(order):
    from apps.polaris.models import PolarisCatalogPart

    brp_part = _part("BRP card", "BRP")
    polaris_part = _part("Polaris card", "POLARIS")
    BrpCatalogPart.objects.create(material_no="001234", part_desc="BRP SEAL")
    PolarisCatalogPart.objects.create(part_number="001234", part_name="POLARIS GASKET")
    _line(order, 1, article="001.234", manufacturer="BRP", name_ru="", name_en="",
          part_type=brp_part)
    _line(order, 2, article="001.234", manufacturer="POLARIS", name_ru="", name_en="",
          part_type=polaris_part)
    _line(order, 3, article="001.234", manufacturer="SPI", name_ru="", name_en="",
          part_type=brp_part)

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert [sheet[f"B{row}"].value for row in (10, 11, 12)] == [
        "001.234", "001.234", "001.234",
    ]
    assert [sheet[f"D{row}"].value for row in (10, 11, 12)] == [
        "BRP SEAL", "POLARIS GASKET", None,
    ]


def test_order_export_does_not_guess_between_normalized_catalog_collisions(order):
    part = _part("Ambiguous BRP card", "BRP")
    BrpCatalogPart.objects.create(material_no="001234", part_desc="FIRST DESCRIPTION")
    BrpCatalogPart.objects.create(material_no="00.1234", part_desc="SECOND DESCRIPTION")
    _line(order, article="00-1234", manufacturer="BRP", name_ru="", name_en="",
          part_type=part)

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert sheet["B10"].value == "00-1234"
    assert sheet["D10"].value is None


def test_order_export_preserves_existing_names_and_leaves_unknown_names_blank(order):
    part = _part("Unmapped catalog card", "BRP")
    BrpCatalogPart.objects.create(material_no="OTHER", part_desc="UNRELATED NAME")
    _line(order, 1, article="001.234", manufacturer="BRP", name_ru="СНИМОК RU",
          name_en="", part_type=part)
    _line(order, 2, article="NO-MATCH", manufacturer="BRP", name_ru="", name_en="",
          part_type=part)

    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    sheet = book["Оригиналы"]
    assert sheet["C10"].value == "СНИМОК RU"
    assert sheet["D10"].value is None
    assert sheet["C11"].value is None
    assert sheet["D11"].value is None
    order.refresh_from_db()
    assert list(order.lines.values_list("name_ru", "name_en")) == [
        ("СНИМОК RU", ""), ("", ""),
    ]


def test_order_export_name_fallback_queries_are_bounded(order, django_assert_num_queries):
    from apps.brp.models import BrpPartLink

    for index in range(20):
        article = f"BATCH-{index:02}"
        part = _part(f"Batch part {index}", "BRP")
        catalog = BrpCatalogPart.objects.create(material_no=article, part_desc=f"PART {index}")
        BrpPartLink.objects.create(
            part=part, brp_part=catalog, usd_rate_used=Decimal("100"),
            markup_percent_used=Decimal("0"),
        )
        _line(order, index + 1, article=article, manufacturer="BRP", name_ru="", name_en="",
              part_type=part)

    # Lines, cards, current customs data, historical versions, and the three
    # catalog families are loaded in batches instead of per-line queries.
    with django_assert_num_queries(7):
        book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    assert len(_article_rows(book["Оригиналы"])) == 20


def test_download_fills_names_without_writing_order_or_customs_data(
    order, client, django_user_model,
):
    from apps.actions.models import PartCustomsInfo
    from apps.brp.models import BrpPartLink

    part = _part("Download card", "BRP")
    catalog = BrpCatalogPart.objects.create(material_no="DOWNLOAD-1", part_desc="OIL SEAL")
    BrpPartLink.objects.create(
        part=part, brp_part=catalog, usd_rate_used=Decimal("100"),
        markup_percent_used=Decimal("0"),
    )
    line = _line(order, article="DOWNLOAD-1", manufacturer="BRP", name_ru="", name_en="",
                 part_type=part)
    order.status = CustomsOrder.Status.FINALIZED
    order.save(update_fields=["status"])
    client.force_login(django_user_model.objects.create_superuser(
        username="customs-export", password="test",
    ))
    customs_count = PartCustomsInfo.objects.count()

    response = client.get(reverse("customs_order_detail", args=[order.pk]), {"export": "1"})

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="customs_order_125.xlsx"'
    sheet = openpyxl.load_workbook(BytesIO(response.content))["Оригиналы"]
    assert sheet["C10"].value == "САЛЬНИК"
    assert sheet["D10"].value == "OIL SEAL"
    line.refresh_from_db()
    order.refresh_from_db()
    assert (line.name_ru, line.name_en) == ("", "")
    assert PartCustomsInfo.objects.count() == customs_count
    assert (order.fx_rate, order.total_quantity, order.total_rub) == (
        Decimal("97.1250"), Decimal("6.500"), Decimal("37498.90"),
    )


def test_order_export_aggregates_identical_snapshots_and_keeps_analog_split(order):
    _line(order)
    _line(
        order, source=CustomsOrderLine.Source.ORDERED, is_ordered=True,
        quantity=Decimal("1.000"), rub_amount=Decimal("995.53"),
    )
    _line(
        order, 2, article="SM-01357", manufacturer="SPI", is_analog=True,
        name_ru="СТАТОР SKI-DOO", name_en="SPI STATOR SKI-DOO", country="",
        quantity=Decimal("1.000"), wholesale_usd=Decimal("203.2600"), occurred_at=_occurred(7),
    )
    _line(
        order, source=CustomsOrderLine.Source.REPAIR, article="SM-09374",
        manufacturer="SPI", is_analog=True, country="", quantity=Decimal("2.000"),
        name_ru="ЩЕКА КОЛЕНЧАТОГО ВАЛА", name_en="SPI PTO CRANK WEB",
        wholesale_usd=Decimal("127.2100"), occurred_at=_occurred(9),
    )
    other_order = CustomsOrder.objects.create(number=126, fx_rate=Decimal("98"))
    _line(other_order, 999, article="NOT-IN-125")

    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    assert book.sheetnames == ["Оригиналы", "Аналоги"]
    originals, analogs = book.worksheets
    assert [originals[f"B{row}"].value for row in _article_rows(originals)] == ["420001234"]
    assert [analogs[f"B{row}"].value for row in _article_rows(analogs)] == [
        "SM-01357", "SM-09374",
    ]
    assert originals["B10"].fill.fgColor.rgb == "FFC6EFCE"
    assert analogs["B10"].fill.fgColor.rgb != "FFC6EFCE"
    assert analogs["B11"].fill.fgColor.rgb != "FFC6EFCE"
    assert analogs["K10"].value == 203.26
    assert analogs["K11"].value == 127.21
    exported_quantity = sum(
        Decimal(str(sheet[f"J{row}"].value))
        for sheet in book for row in _article_rows(sheet)
    )
    assert exported_quantity == Decimal("6.500")
    assert originals["J10"].value == 3.5
    assert sum(len(_article_rows(sheet)) for sheet in book) == 3


def test_order_export_splits_same_article_when_unit_price_conflicts(order):
    _line(order, wholesale_usd=Decimal("10.2500"), quantity=Decimal("2"))
    _line(order, 2, wholesale_usd=Decimal("12.5000"), quantity=Decimal("3"))

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert [sheet[f"J{row}"].value for row in _article_rows(sheet)] == [2, 3]
    assert [sheet[f"K{row}"].value for row in _article_rows(sheet)] == [10.25, 12.5]


def test_order_export_promotes_one_unambiguous_populated_metadata_value(order):
    _line(order, application_area="", gross_weight_kg=None, net_weight_kg=None)
    _line(
        order, 2, application_area="КВАДРОЦИКЛ", gross_weight_kg=Decimal("0.456"),
        net_weight_kg=Decimal("0.321"),
    )

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert _article_rows(sheet) == [10]
    assert sheet["J10"].value == 5
    assert sheet["G10"].value == 0.456
    assert sheet["H10"].value == 0.321
    assert sheet["M10"].value == "КВАДРОЦИКЛ"


def test_order_export_does_not_silently_merge_conflicting_metadata(order):
    _line(order, country="CANADA", application_area="СНЕГОХОД")
    _line(order, 2, country="JAPAN", application_area="КВАДРОЦИКЛ")
    _line(order, 3, country="", application_area="")

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert [sheet[f"F{row}"].value for row in _article_rows(sheet)] == ["CANADA", "JAPAN", None]
    assert [sheet[f"M{row}"].value for row in _article_rows(sheet)] == [
        "СНЕГОХОД", "КВАДРОЦИКЛ", None,
    ]


def test_order_export_never_aggregates_lines_without_a_canonical_article(order):
    _line(order, article="", quantity=Decimal("1"))
    _line(order, 2, article="", quantity=Decimal("2"))

    sheet = openpyxl.load_workbook(export_customs_order_xlsx(order))["Оригиналы"]
    assert [sheet[f"J{row}"].value for row in (10, 11)] == [1, 2]


def test_order_export_orders_each_section_by_frozen_history_chronology(order):
    # Insert deliberately out of source-operation order.  Query/PK order must
    # have no effect on the XLSX sequence.
    _line(order, 30, article="LATE", occurred_at=_occurred(12))
    _line(order, 10, article="EARLY", occurred_at=_occurred(7))
    _line(order, 20, article="MIDDLE", occurred_at=_occurred(9))
    _line(order, 40, article="ANALOG-LATE", is_analog=True, occurred_at=_occurred(12))
    _line(order, 11, article="ANALOG-EARLY", is_analog=True, occurred_at=_occurred(7))

    assert [row[0] for row in _export_rows(order)] == ["EARLY", "MIDDLE", "LATE"]
    assert [row[0] for row in _export_rows(order, "Аналоги")] == [
        "ANALOG-EARLY", "ANALOG-LATE",
    ]


def test_order_export_aggregated_row_uses_earliest_member_chronology(order):
    _line(order, 30, article="MERGED", quantity=Decimal("2"), occurred_at=_occurred(12))
    _line(order, 10, article="MERGED", quantity=Decimal("3"), occurred_at=_occurred(7))
    _line(order, 20, article="BETWEEN", occurred_at=_occurred(9))

    rows = _export_rows(order)
    assert [row[0] for row in rows] == ["MERGED", "BETWEEN"]
    assert rows[0][8] == 5


def test_order_export_merged_article_sorts_by_its_first_operation_date(order):
    # Y is inserted first, so a PK/first-seen order would put it on top.
    _line(order, 30, article="Y", quantity=Decimal("1"), occurred_at=_occurred(8))
    _line(order, 20, article="X", quantity=Decimal("2"), occurred_at=_occurred(10))
    _line(order, 10, article="X", quantity=Decimal("1"), occurred_at=_occurred(7))

    rows = _export_rows(order)
    assert [(row[0], row[8]) for row in rows] == [("X", 3), ("Y", 1)]


def test_order_export_keeps_price_and_metadata_conflicts_in_chronological_order(order):
    _line(
        order, 30, article="CONFLICT", country="JAPAN", wholesale_usd=Decimal("12"),
        occurred_at=_occurred(12),
    )
    _line(order, 20, article="BETWEEN", occurred_at=_occurred(9))
    _line(
        order, 10, article="CONFLICT", country="CANADA", wholesale_usd=Decimal("10"),
        occurred_at=_occurred(7),
    )

    rows = _export_rows(order)
    assert [(row[0], row[4], row[9]) for row in rows] == [
        ("CONFLICT", "CANADA", 10),
        ("BETWEEN", "CANADA", 10.25),
        ("CONFLICT", "JAPAN", 12),
    ]


def test_order_export_different_price_rows_each_take_their_own_first_date(order):
    _line(order, 40, article="PRICED", wholesale_usd=Decimal("12"), occurred_at=_occurred(8))
    _line(order, 30, article="PRICED", wholesale_usd=Decimal("10"), occurred_at=_occurred(11))
    _line(order, 20, article="OTHER", occurred_at=_occurred(9))
    _line(order, 10, article="PRICED", wholesale_usd=Decimal("12"), occurred_at=_occurred(12))

    rows = _export_rows(order)
    assert [(row[0], row[8], row[9]) for row in rows] == [
        ("PRICED", 5, 12), ("OTHER", 2.5, 10.25), ("PRICED", 2.5, 10),
    ]


def test_order_export_same_price_metadata_conflict_rows_keep_chronology(order):
    _line(order, 30, article="META", country="JAPAN", occurred_at=_occurred(7))
    _line(order, 20, article="OTHER", occurred_at=_occurred(8))
    _line(order, 10, article="META", country="CANADA", occurred_at=_occurred(9))

    rows = _export_rows(order)
    assert [(row[0], row[4]) for row in rows] == [
        ("META", "JAPAN"), ("OTHER", "CANADA"), ("META", "CANADA"),
    ]


def test_order_export_breaks_same_datetime_ties_like_customs_history(order):
    occurred_at = _occurred(7)
    _line(order, 30, article="SALE-LATE-ID", occurred_at=occurred_at)
    _line(order, 10, article="SALE-EARLY-ID", occurred_at=occurred_at)
    _line(
        order, 1, source=CustomsOrderLine.Source.REPAIR, article="REPAIR-FIRST",
        occurred_at=occurred_at,
    )

    assert [row[0] for row in _export_rows(order)] == [
        "REPAIR-FIRST", "SALE-EARLY-ID", "SALE-LATE-ID",
    ]


@pytest.mark.parametrize("classification", [None, False, True])
def test_order_export_always_contains_both_sheets_even_when_one_is_empty(order, classification):
    if classification is not None:
        _line(order, is_analog=classification)
    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    assert book.sheetnames == ["Оригиналы", "Аналоги"]
    assert len(_article_rows(book["Оригиналы"])) == int(classification is False)
    assert len(_article_rows(book["Аналоги"])) == int(classification is True)
    for sheet in book:
        assert sheet["I150"].value is None
        assert sheet["B12"].value is None


def test_order_export_uses_all_frozen_fields_after_catalog_and_fx_change(
    order, django_assert_num_queries,
):
    _line(
        order, gross_weight_kg=Decimal("0.456"), net_weight_kg=Decimal("0.321"),
        country="JAPAN", application_area="КВАДРОЦИКЛ",
    )
    BrpCatalogPart.objects.create(
        material_no="420001234", part_desc="CHANGED CATALOG NAME",
        wholesale_price_usd=Decimal("999.99"), retail_price_usd=Decimal("1999"),
    )
    settings = ValuationSettings.get()
    settings.current_usd_rate = Decimal("150")
    settings.save(update_fields=["current_usd_rate"])

    # Source ids need not resolve: the only data read is the saved order lines.
    with django_assert_num_queries(1):
        book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    sheet = book["Оригиналы"]
    assert [sheet[f"{column}10"].value for column in "BCDEFGH"] == [
        "420001234", "ПОДШИПНИК", "BALL BEARING", "BRP", "JAPAN", 0.456, 0.321,
    ]
    assert sheet["J10"].value == 2.5
    assert sheet["K10"].value == 10.25
    assert sheet["M10"].value == "КВАДРОЦИКЛ"
    assert sheet["I10"].value == "=J10*G10"
    assert sheet["L10"].value == "=K10*J10"
    assert sheet["I150"].value == "=SUM(I10:I149)"
    order.refresh_from_db()
    assert order.fx_rate == Decimal("97.1250")
    assert order.total_rub == Decimal("37498.90")


def test_both_sheets_keep_template_headers_dimensions_drawings_and_data_style(order):
    long_name = "КОМПЛЕКТ ОБСЛУЖИВАНИЯ СЦЕПЛЕНИЯ EDRIVE II"
    _line(order, name_ru=long_name)
    _line(order, 2, is_analog=True, name_ru=long_name)
    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    template = openpyxl.load_workbook(TEMPLATE_PATH).active
    for sheet in book:
        for row in template.iter_rows(min_row=1, max_row=9):
            for original in row:
                cell = sheet[original.coordinate]
                expected = original.value
                if original.column == 9 and 7 <= original.row <= 9:
                    expected = (
                        f'=IF(AND(ISNUMBER(J{original.row}),ISNUMBER(G{original.row})),'
                        f'J{original.row}*G{original.row},"")'
                    )
                elif original.column == 12 and 7 <= original.row <= 9:
                    expected = (
                        f'=IF(AND(ISNUMBER(K{original.row}),ISNUMBER(J{original.row})),'
                        f'K{original.row}*J{original.row},"")'
                    )
                assert cell.value == expected
                # This supplier template has no default style. openpyxl
                # recreates style registries on every load, so object and
                # style-id equality across workbooks are not meaningful.
                assert str(cell.font) == str(original.font)
                assert str(cell.fill) == str(original.fill)
                assert str(cell.border) == str(original.border)
                assert cell.number_format == original.number_format
        assert str(sheet.merged_cells) == str(template.merged_cells)
        assert sheet.sheet_format == template.sheet_format
        assert sheet.page_setup == template.page_setup
        assert sheet.page_margins == template.page_margins
        assert len(sheet._images) == len(template._images)
        assert {key: value.width for key, value in sheet.column_dimensions.items()} == {
            key: value.width for key, value in template.column_dimensions.items()
        }
        for column in "ABCDEFGHIJKLM":
            cell = sheet[f"{column}10"]
            assert (cell.font.name, cell.font.sz) == ("Arial", 12)
            assert cell.alignment.wrap_text is True
            # Excel's false boolean is normalized by openpyxl to None on load.
            assert cell.alignment.shrink_to_fit is not True
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
        assert sheet.row_dimensions[10].height == 30
        assert sheet["F10"].value == "CANADA"
        for column in "GHI":
            assert sheet[f"{column}10"].number_format == "0.00"
        assert sheet["G10"].value is None
        assert sheet["H10"].value is None
        assert sheet["I10"].value is None
        assert sheet["L10"].value == "=K10*J10"


def test_large_order_preserves_each_analog_row_and_moves_total_below_data(order):
    for source_id in range(1, 143):
        _line(order, source_id, article=f"SPI-{source_id}", is_analog=True)
    book = openpyxl.load_workbook(export_customs_order_xlsx(order))
    sheet = book["Аналоги"]
    assert len(_article_rows(sheet)) == 142
    assert sheet["B151"].value == "SPI-142"
    assert sheet["I151"].value is None
    assert sheet["L151"].value == "=K151*J151"
    assert sheet["I152"].value is None
    assert "F152:H152" in sheet.merged_cells
    assert "F150:H150" not in sheet.merged_cells
    assert sheet["H150"].value is None
    assert sheet["C151"].font.sz == 12
    assert sheet["G151"].number_format == "0.00"
    assert _article_rows(book["Оригиналы"]) == []
