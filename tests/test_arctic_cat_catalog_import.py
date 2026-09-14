"""Arctic Cat adapter: source facts only, through the normal import workflow."""

from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import Workbook

from apps.brp.models import BrpPricingSettings
from apps.brp.pricing import customer_price_rub
from apps.catalog.models import PartNumber, PartType
from apps.catalog.price_audit import EXACT_MATCH, audit_prices
from apps.catalog.services import (
    certify_valid_manual_price_exception,
    create_manual_part,
    plan_linked_part_price_refresh,
    refresh_linked_part_prices,
)
from apps.catalog_import.models import ArcticCatCatalogPart, CatalogImportBatch
from apps.catalog_import.services import CatalogImportError, apply_batch, run_check, save_upload
from apps.inventory.models import PartItem, StockBalance, StockLot, StockMovement
from apps.warehouse.models import ValuationSettings

HEADERS = ["P/N", "Description", "Pkg Qty", "DEALER PRICE"]
PASSWORD = "arctic-test-password"


def _workbook(rows, *, sheet="usprice", headers=HEADERS):
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append(headers)
    for row in rows:
        page.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _upload(rows, **kwargs):
    return SimpleUploadedFile(
        "ARCTIC CAT DEALER 2026.xlsx",
        _workbook(rows, **kwargs),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


RATE = Decimal("105")
MARKUP = Decimal("40")


@pytest.fixture(autouse=True)
def rates(db):
    valuation = ValuationSettings.get()
    valuation.current_usd_rate = RATE
    valuation.save()
    markup = BrpPricingSettings.get()
    markup.brp_markup_percent = MARKUP
    markup.save()


def _canonical(usd):
    """Expected price from the ONE project formula, not a copy of its arithmetic."""
    return customer_price_rub(Decimal(usd), RATE, MARKUP)


@pytest.fixture
def boss(client, django_user_model):
    user = django_user_model.objects.create_superuser("arctic-boss", password=PASSWORD)
    client.force_login(user)
    return client


def _send(client, rows, **kwargs):
    return client.post(
        reverse("catalog_import_upload"),
        {"catalog": "arctic_cat", "workbook": _upload(rows, **kwargs)},
        follow=True,
    )


def _apply(client, rows, **kwargs):
    _send(client, rows, **kwargs)
    batch = CatalogImportBatch.objects.latest("pk")
    client.post(reverse("catalog_import_apply", args=[batch.pk]), follow=True)
    batch.refresh_from_db()
    assert batch.status == CatalogImportBatch.Status.APPLIED, batch.error_text
    return batch


def _stock_snapshot():
    return (
        StockLot.objects.count(),
        StockBalance.objects.count(),
        PartItem.objects.count(),
        StockMovement.objects.count(),
    )


def test_valid_usprice_is_dry_run_then_creates_private_source_facts_only(boss):
    before = _stock_snapshot()
    response = _send(boss, [["0101-045", "KEY, TORQUE BRK", "5", "1.61"]])

    batch = CatalogImportBatch.objects.get()
    assert response.status_code == 200
    assert batch.catalog == CatalogImportBatch.Catalog.ARCTIC_CAT
    assert batch.status == CatalogImportBatch.Status.CHECKED
    assert batch.summary["sheet"] == "usprice"
    assert batch.summary["new"] == 1
    assert PartType.objects.count() == 0

    boss.post(reverse("catalog_import_apply", args=[batch.pk]), follow=True)
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert entry.supplier_article == "0101-045"
    assert entry.source_description == "KEY, TORQUE BRK"
    assert entry.package_quantity == "5"
    assert entry.dealer_price_usd == Decimal("1.61")
    assert entry.part.recommended_price == _canonical("1.61")
    assert entry.part.certified_price_rub == entry.part.recommended_price
    assert entry.part.price_provenance == PartType.PriceProvenance.FORMULA_CERTIFIED
    assert not entry.part.is_public
    assert entry.part.numbers.get(kind=PartNumber.Kind.ARTICLE).value == "0101-045"
    assert _stock_snapshot() == before


def test_numeric_cell_with_safe_excel_format_keeps_leading_zero_article(boss):
    book = Workbook()
    page = book.active
    page.title = "usprice"
    page.append(HEADERS)
    page.append([101045, "KEY", "5", "1.61"])
    page.cell(2, 1).number_format = "0000-000"
    output = BytesIO()
    book.save(output)
    boss.post(
        reverse("catalog_import_upload"),
        {
            "catalog": "arctic_cat",
            "workbook": SimpleUploadedFile("arctic.xlsx", output.getvalue()),
        },
        follow=True,
    )
    batch = CatalogImportBatch.objects.get()
    boss.post(reverse("catalog_import_apply", args=[batch.pk]), follow=True)
    assert ArcticCatCatalogPart.objects.get().supplier_article == "0101-045"


@pytest.mark.parametrize(
    ("price", "state", "counter"),
    [("0", "zero", "zero_price_rows"), ("", "blank", "blank_price_rows")],
)
def test_zero_or_blank_dealer_price_is_unavailable_not_a_customer_price(
    boss, price, state, counter
):
    batch = _apply(boss, [["0101-045", "KEY", "5", price]])
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert batch.apply_summary[counter] == 1
    assert entry.dealer_price_usd is None
    assert entry.dealer_price_state == state
    # No fabricated 0 ₽: the card exists, but is honestly unpriced.
    assert entry.part.recommended_price is None
    assert entry.part.certified_price_rub is None
    assert entry.part.price_provenance == PartType.PriceProvenance.SOURCE_MISSING
    assert not entry.part.is_public


def test_unavailable_reimport_never_erases_a_known_positive_supplier_price(boss):
    _apply(boss, [["0101-045", "KEY", "5", "1.61"]])
    _apply(boss, [["0101-045", "KEY", "5", "0"]])
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert entry.dealer_price_usd == Decimal("1.61")
    assert entry.dealer_price_state == "known"
    assert entry.part.recommended_price == _canonical("1.61")
    assert entry.part.price_provenance == PartType.PriceProvenance.FORMULA_CERTIFIED


def test_exact_replacement_is_skipped_without_creating_any_record(boss):
    _send(boss, [["0101-058", "R/B 0101-055", "1", "0"]])
    batch = CatalogImportBatch.objects.get()
    assert batch.summary["skipped_rb"] == 1
    assert batch.summary["valid"] == 0
    assert not ArcticCatCatalogPart.objects.exists()


def test_exact_rb_rows_create_no_card_number_price_or_relation_next_to_real_rows(boss):
    before = _stock_snapshot()
    batch = _apply(
        boss,
        [
            ["0101-045", "KEY", "5", "1.61"],
            ["0101-058", "R/B 0101-055", "10", "348.58"],
            ["0409-200", "r/b H680507", "abc", ""],
        ],
    )
    assert batch.summary["skipped_rb"] == 2
    assert batch.apply_summary["skipped_rb"] == 2
    assert batch.apply_summary["created_parts"] == 1
    # A skipped row contributes to no price or package counter and no warning.
    assert batch.summary["positive_price_rows"] == 1
    assert batch.summary["blank_price_rows"] == 0
    assert batch.summary["invalid_package_quantity_rows"] == 0
    assert batch.summary["warnings"] == 0
    assert PartType.objects.count() == 1
    assert ArcticCatCatalogPart.objects.get().supplier_article == "0101-045"
    for number in ("0101058", "0101055", "0409200", "H680507"):
        assert not PartNumber.objects.filter(normalized_value=number).exists()
    assert not ArcticCatCatalogPart.objects.exclude(replacement_article="").exists()
    assert _stock_snapshot() == before


def test_malformed_replacement_is_a_warning_and_stays_plain_description(boss):
    response = _send(boss, [["0101-058", "R/B 0101-055 EXTRA", "1", "0"]])
    batch = CatalogImportBatch.objects.get()
    assert response.status_code == 200
    assert batch.status == CatalogImportBatch.Status.CHECKED
    assert batch.summary["warnings"] == 1
    boss.post(reverse("catalog_import_apply", args=[batch.pk]), follow=True)
    entry = ArcticCatCatalogPart.objects.get()
    assert entry.source_description == "R/B 0101-055 EXTRA"
    assert entry.replacement_article == ""


def test_real_style_alphanumeric_replacement_is_skipped(boss):
    _send(boss, [["0409-200", "R/B H680507", "1", "348.58"]])
    assert CatalogImportBatch.objects.get().summary["skipped_rb"] == 1
    assert not ArcticCatCatalogPart.objects.exists()


def test_invalid_package_quantity_is_a_warning_and_never_stock(boss):
    before = _stock_snapshot()
    batch = _apply(boss, [["0101-045", "KEY", "0", "1.61"]])
    entry = ArcticCatCatalogPart.objects.get()
    assert batch.apply_summary["invalid_package_quantity_rows"] == 1
    assert entry.package_quantity == "0"
    assert _stock_snapshot() == before


def test_arctic_preview_explains_price_policy_and_package_warning_correctly(boss):
    response = _send(boss, [["0101-045", "KEY", "0", "1.61"]])

    batch = CatalogImportBatch.objects.get()
    body = response.content.decode()
    assert batch.status == CatalogImportBatch.Status.CHECKED
    assert batch.summary["invalid_package_quantity_rows"] == 1
    assert "цена одной детали в USD" in body
    assert "централизованным курсу USD" in body
    assert "помечает её как подтверждённую формулой" in body
    assert "Предупреждение о некорректном" in body
    assert "Pkg Qty не исключает строку" in body
    assert "ещё нет утверждённой политики цены клиента" not in body
    assert "Эти строки не применяются" not in body


def test_duplicates_are_reported_and_conflicts_are_never_arbitrarily_applied(boss):
    _send(
        boss,
        [
            ["0101-045", "KEY A", "5", "1.61"],
            ["0101-045", "KEY B", "5", "2.00"],
        ],
    )
    batch = CatalogImportBatch.objects.get()
    assert batch.summary["duplicate_part_numbers"] == 1
    assert batch.summary["errors"] == 1
    boss.post(reverse("catalog_import_apply", args=[batch.pk]), follow=True)
    assert ArcticCatCatalogPart.objects.get().source_description == "KEY A"


def test_changes_are_classified_and_repeated_file_is_idempotent(boss):
    _apply(boss, [["0101-045", "OLD", "5", "1.61"]])
    changed = _apply(boss, [["0101-045", "NEW", "10", "2.00"]])
    assert changed.apply_summary["existing"] == 1
    assert changed.apply_summary["description_changed"] == 1
    assert changed.apply_summary["package_quantity_changed"] == 1
    assert changed.apply_summary["price_changed"] == 1
    assert changed.apply_summary["repriced_parts"] == 1
    part = ArcticCatCatalogPart.objects.select_related("part").get().part
    assert part.recommended_price == _canonical("2.00")

    repeat = _apply(boss, [["0101-045", "NEW", "10", "2.00"]])
    assert repeat.apply_summary["unchanged"] == 1
    assert repeat.apply_summary["updated_parts"] == 0
    assert repeat.apply_summary["repriced_parts"] == 0
    assert ArcticCatCatalogPart.objects.count() == 1


def test_rb_row_for_an_existing_article_leaves_the_existing_card_untouched(boss):
    _apply(boss, [["0101-045", "KEY", "5", "1.61"]])
    _send(boss, [["0101-045", "R/B 0101-055", "10", "2.00"]])
    assert CatalogImportBatch.objects.latest("pk").summary["skipped_rb"] == 1
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert (entry.source_description, entry.dealer_price_usd) == ("KEY", Decimal("1.61"))
    assert entry.replacement_article == ""


def test_same_article_in_another_catalog_is_not_merged(boss):
    manual = create_manual_part(
        name="Manual Arctic-like article",
        article="0101-045",
        manufacturer_name="Another manufacturer",
    )
    _apply(boss, [["0101-045", "Arctic source", "1", "1.61"]])
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert entry.part_id != manual.pk
    assert PartType.objects.filter(numbers__normalized_value="0101045").distinct().count() == 2


def test_search_finds_source_article_by_exact_normalized_partial_and_description(boss):
    _apply(boss, [["0101-045", "TORQUE KEY", "1", "1.61"]])
    for query in ("0101-045", "0101045", "0101", "TORQUE"):
        response = boss.get(reverse("part_search"), {"q": query})
        assert "TORQUE KEY" in response.content.decode()


def test_stale_dry_run_is_rejected(boss):
    _send(boss, [["0101-045", "KEY", "1", "1.61"]])
    batch = CatalogImportBatch.objects.get()
    _apply(boss, [["0101-046", "OTHER", "1", "2.00"]])
    with pytest.raises(CatalogImportError, match="STALE_DRY_RUN"):
        apply_batch(batch)


def test_wrong_schema_can_be_inspected_without_any_catalog_write(boss):
    _send(boss, [["0101-045", "KEY"]], sheet="Price", headers=["Article", "Name"])
    batch = CatalogImportBatch.objects.get()
    assert batch.status == CatalogImportBatch.Status.CHECK_FAILED
    response = boss.get(reverse("catalog_import_inspect", args=[batch.pk]))
    body = response.content.decode()
    assert response.status_code == 200
    assert "Price" in body
    assert "P/N" in body
    assert "Отсутствуют обязательные заголовки" in body
    assert not ArcticCatCatalogPart.objects.exists()


def test_broken_workbook_is_rejected_without_source_rows(boss):
    response = boss.post(
        reverse("catalog_import_upload"),
        {
            "catalog": "arctic_cat",
            "workbook": SimpleUploadedFile("broken.xlsx", b"not an xlsx archive"),
        },
        follow=True,
    )
    assert "не читается" in response.content.decode().lower()
    assert not ArcticCatCatalogPart.objects.exists()


def test_upload_requires_same_permission_as_existing_catalog_import(client, django_user_model):
    user = django_user_model.objects.create_user("seller", password=PASSWORD)
    client.force_login(user)
    response = _send(client, [["0101-045", "KEY", "1", "1.61"]])
    assert response.status_code == 403
    assert not CatalogImportBatch.objects.exists()


def test_private_upload_and_service_apply_have_no_public_url_or_direct_stock_write(
    boss, settings, tmp_path
):
    settings.PRIVATE_MEDIA_ROOT = str(tmp_path / "private")
    upload = _upload([["0101-045", "KEY", "1", "1.61"]])
    batch = run_check(save_upload(upload, catalog="arctic_cat"))
    assert batch.stored_path and "/" not in batch.stored_path
    before = _stock_snapshot()
    apply_batch(batch)
    assert _stock_snapshot() == before


# --- Owner-confirmed pricing: DEALER PRICE is USD per ONE part ------------------------


def test_dealer_price_is_per_unit_and_package_quantity_never_scales_it(boss):
    _apply(boss, [["0101-045", "WASHER", "10", "1.08"]])
    entry = ArcticCatCatalogPart.objects.select_related("part").get()
    assert entry.dealer_price_usd == Decimal("1.08")
    assert entry.package_quantity == "10"
    assert entry.part.recommended_price == _canonical("1.08")
    assert entry.part.recommended_price != _canonical("10.80")
    assert entry.part.recommended_price != _canonical("0.108")


def test_same_dealer_price_with_different_package_quantities_gives_same_unit_price(boss):
    _apply(
        boss,
        [
            ["0101-001", "A", "1", "7.39"],
            ["0101-002", "B", "10", "7.39"],
            ["0101-003", "C", "250", "7.39"],
            ["0101-004", "D", "", "7.39"],
            ["0101-005", "E", "bag", "7.39"],
        ],
    )
    prices = set(PartType.objects.values_list("recommended_price", flat=True))
    assert prices == {_canonical("7.39")}
    assert set(PartType.objects.values_list("price_provenance", flat=True)) == {
        PartType.PriceProvenance.FORMULA_CERTIFIED
    }


@pytest.mark.parametrize(("usd", "rub"), [("7.39", "1086"), ("9.03", "1327"), ("99.99", "14699")])
def test_rounding_is_the_canonical_whole_rouble_half_up(boss, usd, rub):
    _apply(boss, [["0101-045", "KEY", "1", usd]])
    assert PartType.objects.get().recommended_price == Decimal(rub)


def test_price_follows_the_shared_settings_not_constants_in_the_importer(boss):
    valuation = ValuationSettings.get()
    valuation.current_usd_rate = Decimal("90")
    valuation.save()
    markup = BrpPricingSettings.get()
    markup.brp_markup_percent = Decimal("25")
    markup.save()
    batch = _apply(boss, [["0101-045", "KEY", "1", "10.00"]])
    assert PartType.objects.get().recommended_price == customer_price_rub(
        Decimal("10.00"), Decimal("90"), Decimal("25")
    )
    assert Decimal(batch.summary["usd_rate"]) == Decimal("90")
    assert Decimal(batch.apply_summary["markup_percent"]) == Decimal("25")


def test_imported_arctic_cards_are_never_published(boss):
    _apply(boss, [["0101-045", "KEY", "1", "1.61"], ["0101-046", "NUT", "1", "0"]])
    assert not PartType.objects.filter(is_public=True).exists()


def test_central_recalculation_keeps_arctic_cards_certified_and_is_idempotent(boss):
    _apply(boss, [["0101-045", "KEY", "5", "1.61"], ["0101-046", "NUT", "5", "0"]])
    kwargs = dict(usd_rate=RATE, brp_markup=MARKUP, polaris_markup=MARKUP)
    plan = plan_linked_part_price_refresh(**kwargs)
    assert plan.arctic_cat_links == 2
    assert plan.parts_to_update == {}
    assert refresh_linked_part_prices(**kwargs) == 0
    priced = PartType.objects.get(numbers__normalized_value="0101045")
    unpriced = PartType.objects.get(numbers__normalized_value="0101046")
    assert priced.price_provenance == PartType.PriceProvenance.FORMULA_CERTIFIED
    assert unpriced.price_provenance == PartType.PriceProvenance.SOURCE_MISSING
    assert unpriced.recommended_price is None

    refresh_linked_part_prices(usd_rate=Decimal("100"), brp_markup=MARKUP, polaris_markup=MARKUP)
    priced.refresh_from_db()
    assert priced.recommended_price == customer_price_rub(Decimal("1.61"), Decimal("100"), MARKUP)
    assert priced.certified_price_rub == priced.recommended_price


def test_manual_exception_on_an_arctic_card_survives_import_and_recalculation(boss):
    _apply(boss, [["0101-045", "KEY", "5", "1.61"]])
    part = PartType.objects.get()
    part.recommended_price = Decimal("45000")
    part.save(update_fields=["recommended_price"])
    certify_valid_manual_price_exception(part)

    _apply(boss, [["0101-045", "KEY", "5", "2.00"]])
    refresh_linked_part_prices(usd_rate=RATE, brp_markup=MARKUP, polaris_markup=MARKUP)
    part.refresh_from_db()
    assert part.recommended_price == Decimal("45000")
    assert part.price_provenance == PartType.PriceProvenance.VALID_MANUAL_EXCEPTION


def test_price_audit_verifies_arctic_cards_against_the_canonical_formula(boss):
    _apply(boss, [["0101-045", "KEY", "5", "1.61"]])
    report = audit_prices(usd_rate=RATE, brp_markup=MARKUP, polaris_markup=MARKUP)
    assert report.by_source["arctic_cat"] == 1
    assert report.by_category[EXACT_MATCH] == 1


def test_apply_query_count_does_not_grow_with_the_number_of_rows(boss, tmp_path):
    from apps.catalog_import.arctic_cat_catalog import apply_file

    def run(size, offset):
        rows = [[f"{offset + i}-{i:03d}", f"PART {i}", "5", "1.61"] for i in range(size)]
        path = tmp_path / f"arctic-{offset}.xlsx"
        path.write_bytes(_workbook(rows))
        with CaptureQueriesContext(connection) as ctx:
            apply_file(path)
        return len(ctx.captured_queries)

    run(1, 100)  # creates the shared category and manufacturer once
    assert run(5, 1000) == run(60, 2000)
